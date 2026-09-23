"""
The agent evaluation set: 30 confirmed findings, and the sheet they are labelled on.

The proposal committed to three measures of the agents -- Macro-F1, a confusion matrix
and the critical miss rate -- against labelled cases. This file holds the cases; the
labels live in `eval/agent_labels.csv`; `eval/agent_eval.py` runs the agents over the
cases and scores them.

Only confirmed violations are here, because nothing else reaches an agent: an
unconfirmed or low-confidence finding stops at the entry gate and goes to the review
queue without a model being called. That gate is tested elsewhere.

The cases vary the four things the Adjudicator is meant to weigh:

    the area          which hazard is present, so which item matters most
    what is missing   one item or several
    identity          known, with 0-3 confirmed prior violations in the last 7 days,
                      or unknown, in which case the history cannot be read at all
    evidence          strong (10 of 10 frames, person clearly detected) or moderate
                      (8 of 10 frames, person detected at 72%)

The labeller sees the area as a supervisor knows it -- its name and what goes on
there -- and never the severity weights, the band table or the rubric's answer. A
label that was read off the rubric would make the evaluation measure the rubric
against itself.

    python eval/agent_cases.py --sheet     # write eval/agent_labels.xlsx to label
    python eval/agent_cases.py --csv       # after labelling: xlsx -> agent_labels.csv
"""

from __future__ import annotations

import argparse
import copy
import csv
import json
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).absolute().parents[1]
sys.path.insert(0, str(ROOT))

from perception.events import DETECTOR_RECALL  # noqa: E402

ACTIONS = ("no_action", "log_only", "warning", "escalation", "stop_work")
SHEET = ROOT / "eval" / "agent_labels.xlsx"
LABELS = ROOT / "eval" / "agent_labels.csv"
WORKER = "W-EVAL"


@dataclass(frozen=True)
class Case:
    case_id: str
    camera: str
    zone: str
    missing: tuple[str, ...]
    identified: bool
    priors: int = 0                 # confirmed, in the window; meaningless if unidentified
    evidence: str = "strong"        # "strong" | "moderate"


CASES = [
    # walkways: transit, head protection first
    Case("A01", "cam_3", "walkway", ("helmet",), False),
    Case("A02", "cam_3", "walkway", ("helmet",), True, 0),
    Case("A03", "cam_3", "walkway", ("helmet",), True, 2),
    Case("A04", "d_view02", "walkway", ("vest",), False),
    Case("A05", "d_view02", "walkway", ("vest",), True, 0),
    Case("A06", "d_view02", "walkway", ("vest",), True, 1),
    Case("A07", "d_view02", "walkway", ("vest",), True, 2),
    Case("A08", "d_view02", "walkway", ("helmet",), False, 0, "moderate"),
    Case("A09", "d_view02", "walkway", ("vest",), False, 0, "moderate"),
    # work area beyond the barriers: manual handling
    Case("B01", "d_view02", "work_area", ("gloves",), False),
    Case("B02", "d_view02", "work_area", ("vest", "gloves"), True, 0),
    Case("B03", "d_view02", "work_area", ("helmet", "vest", "gloves"), False),
    # assembly line: sheet stock
    Case("C01", "cam_1", "assembly_line", ("gloves",), False),
    Case("C02", "cam_1", "assembly_line", ("gloves",), True, 1),
    Case("C03", "cam_1", "assembly_line", ("gloves",), True, 3),
    Case("C04", "cam_1", "assembly_line", ("helmet", "gloves"), False, 0, "moderate"),
    # grinding station: abrasive wheel running
    Case("D01", "cam_3", "grinding_station", ("goggles",), False),
    Case("D02", "cam_3", "grinding_station", ("gloves",), True, 0),
    Case("D03", "cam_3", "grinding_station", ("gloves", "goggles"), False),
    Case("D04", "cam_3", "grinding_station", ("goggles",), True, 2),
    Case("D05", "cam_3", "grinding_station", ("helmet", "gloves", "goggles"), False),
    # welding bay: hot work
    Case("E01", "cam_3", "welding_bay", ("vest",), False),
    Case("E02", "cam_3", "welding_bay", ("vest",), True, 1),
    Case("E03", "cam_3", "welding_bay", ("goggles",), False),
    Case("E04", "cam_3", "welding_bay", ("goggles",), False, 0, "moderate"),
    Case("E05", "cam_3", "welding_bay", ("gloves", "goggles"), True, 1),
    # loading dock: vehicles moving
    Case("F01", "cam_3", "loading_dock", ("vest",), False),
    Case("F02", "cam_3", "loading_dock", ("boots",), True, 0),
    Case("F03", "cam_3", "loading_dock", ("vest", "boots"), False),
    Case("F04", "cam_3", "loading_dock", ("helmet",), True, 1),
]


def _zone(camera: str, zone: str) -> dict:
    zones = json.loads((ROOT / "config" / "zones.json").read_text(encoding="utf-8"))
    for z in zones["cameras"][camera]["zones"]:
        if z["name"] == zone:
            return z
    raise KeyError(f"{camera}/{zone}")


def build_event(case: Case, now: datetime | None = None) -> dict:
    """A ViolationEvent for one case, shaped exactly like the pipeline's output."""
    now = now or datetime.now().astimezone()
    z = _zone(case.camera, case.zone)
    strong = case.evidence == "strong"
    present = {item: 0.88 for item in z["required_ppe"] if item not in case.missing}
    return {
        "event_id": f"evt_eval_{case.case_id}",
        "captured_at": (now - timedelta(hours=1)).isoformat(timespec="seconds"),
        "camera_id": case.camera,
        "zone": {"name": z["name"], "label": z["label"],
                 "required_ppe": list(z["required_ppe"]),
                 "severity_multiplier": z.get("severity_multiplier", 1.0)},
        "subject": {"track_id": 1, "bbox": [500, 280, 700, 700],
                    "detection_confidence": 0.91 if strong else 0.72,
                    "height_px": 420, "worker_ref": WORKER if case.identified else None},
        "ppe": {"present": present, "missing": list(case.missing), "indeterminate": []},
        "status": "violation",
        "confirmation": {"frames_observed": 10, "frames_missing": 10 if strong else 8,
                         "window_seconds": 2.0, "rule": "8_of_10", "confirmed": True},
        "evidence": {"crop_url": None, "frame_url": None,
                     "crop_path": None, "frame_path": None},
        "detector": {"weights": "ppe-presence-s960-3/weights/best.pt", "imgsz": 960,
                     "recall": dict(DETECTOR_RECALL)},
        "reasons": [],
        "schema_version": "1.0",
    }


def seed_history(db, case: Case, now: datetime | None = None) -> None:
    """Put the case's worker on the roster with `priors` confirmed violations, each
    bound by a supervisor, inside the 7-day window -- the only history that counts."""
    from app.db import Worker

    if not case.identified:
        return
    now = now or datetime.now().astimezone()
    db.add_worker(Worker(WORKER, "Evaluation worker"))
    for i in range(case.priors):
        prior = build_event(case, now - timedelta(days=i + 1))
        prior["event_id"] = f"evt_eval_{case.case_id}_prior{i}"
        db.record_event(prior)
        db.attach_identity(prior["event_id"], WORKER, "supervisor:eval")


def rubric_action(case: Case) -> str:
    """What the deterministic rubric alone would do: the band of the final score."""
    from agents.tools import final_severity

    z = _zone(case.camera, case.zone)
    score = final_severity(case.zone, list(case.missing),
                           case.priors if case.identified else 0, case.camera,
                           required=z["required_ppe"])
    return "no_action" if score.band == "compliant" else score.band


def describe(case: Case) -> dict:
    """The case as the labeller sees it: facts, never scores."""
    z = _zone(case.camera, case.zone)
    return {
        "case": case.case_id,
        "area": z["label"],
        "what happens there": z.get("notes", ""),
        "required PPE": ", ".join(z["required_ppe"]),
        "MISSING": ", ".join(case.missing),
        "worker identified?": "yes" if case.identified else "NO",
        "confirmed prior violations, last 7 days": (
            str(case.priors) if case.identified
            else "unknown (not identified, history cannot be read)"),
        "evidence": ("strong: missing in 10 of 10 frames, person clearly detected"
                     if case.evidence == "strong" else
                     "moderate: missing in 8 of 10 frames, person detected at 72%"),
    }


GUIDE = [
    ("What this is", "30 confirmed PPE findings. For each one, choose what the site "
                     "should do, as the safety supervisor would. Your answers are the "
                     "ground truth the AI agents are scored against."),
    ("How to judge", "Use only the facts on the Cases tab and your own judgement. Do "
                     "not look up the scoring rubric or the severity weights, and do "
                     "not try to guess what the system would say. A label copied from "
                     "the rubric makes the evaluation measure the rubric against "
                     "itself."),
    ("Unknown identity", "If the worker is not identified, you cannot see their "
                         "history. Decide as you would in real life with that "
                         "uncertainty. It is neither a clean record nor a bad one."),
    ("Evidence", "Moderate evidence means the finding is real enough to act on but "
                 "not strongly supported. Decide whether that changes your action."),
    ("no_action", "Nothing is issued and nothing is recorded as a finding."),
    ("log_only", "Recorded for the record. No notice is raised."),
    ("warning", "The supervisor has a word with the worker."),
    ("escalation", "Formal. The supervisor must respond, and the case needs a "
                   "signature."),
    ("stop_work", "Halt the activity immediately."),
    ("Afterwards", "Save the file and say so. Optional: if a second person labels a "
                   "copy independently, the report can state how often two humans "
                   "agree, which is the ceiling any system can be measured against."),
]


def write_sheet(path: Path = SHEET) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    rows = [describe(c) for c in CASES]
    header = list(rows[0]) + ["YOUR ACTION"]
    ws.append(header)
    for r in rows:
        ws.append(list(r.values()) + [""])

    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font, cell.alignment = bold, wrap
    widths = [7, 26, 44, 26, 22, 12, 26, 34, 16]
    for col, width in zip("ABCDEFGHI", widths):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
        row[-1].fill = fill
    ws.freeze_panes = "B2"

    choice = DataValidation(type="list", formula1='"' + ",".join(ACTIONS) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Not an action",
                            error="Choose one of: " + ", ".join(ACTIONS))
    ws.add_data_validation(choice)
    choice.add(f"I2:I{len(rows) + 1}")

    guide = wb.create_sheet("How to label", 0)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 100
    for term, text in GUIDE:
        guide.append([term, text])
    for row in guide.iter_rows():
        row[0].font = bold
        for cell in row:
            cell.alignment = wrap
    wb.save(path)
    return path


def sheet_to_csv(sheet: Path = SHEET, out: Path = LABELS) -> Path:
    """Read the labelled sheet and write the labels as a diffable CSV."""
    from openpyxl import load_workbook

    ws = load_workbook(sheet, read_only=True)["Cases"]
    rows = list(ws.iter_rows(values_only=True))
    col = rows[0].index("YOUR ACTION")
    labels, blank = [], []
    for row in rows[1:]:
        if not row[0]:
            continue
        action = (row[col] or "").strip()
        if action not in ACTIONS:
            blank.append(row[0])
        labels.append((row[0], action))
    if blank:
        sys.exit(f"not labelled, or not one of {ACTIONS}: {', '.join(blank)}")
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["case", "label"])
        w.writerows(labels)
    return out


def load_labels(path: Path = LABELS) -> dict[str, str]:
    with open(path, encoding="utf-8") as f:
        return {r["case"]: r["label"] for r in csv.DictReader(f)}


def cases_by_id() -> dict[str, Case]:
    return {c.case_id: copy.copy(c) for c in CASES}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sheet", action="store_true", help="write the sheet to label")
    ap.add_argument("--csv", action="store_true", help="labelled sheet -> CSV")
    args = ap.parse_args()
    if args.sheet:
        print(f"wrote {write_sheet()}")
    elif args.csv:
        print(f"wrote {sheet_to_csv()}")
    else:
        for c in CASES:
            print(c.case_id, describe(c)["area"], c.missing,
                  f"priors={c.priors}" if c.identified else "unidentified",
                  c.evidence, "| rubric:", rubric_action(c))
